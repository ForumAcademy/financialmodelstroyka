"""
Строит полную карту исходного Excel → новая модель.
Результат:
  legacy/legacy_values_map.csv   — КАЖДАЯ ячейка с числом/датой исходника (ничего не теряется)
  legacy/legacy_formulas_map.csv — КАЖДЫЙ уникальный шаблон формулы (по строкам, R1C1-нормализация)
  legacy/legacy_text_cells.csv   — текстовые ячейки (подписи, комментарии автора)
Запуск: python scripts/build_legacy_map.py  (нужен openpyxl, pyyaml)
"""
import csv, re, sys
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
import yaml

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "legacy" / "Кальк Саевой привязка КОД.xlsx"

wb = openpyxl.load_workbook(XLSX)
wv = openpyxl.load_workbook(XLSX, data_only=True)

capex = yaml.safe_load(open(ROOT / "data" / "capex_items.yaml"))["items"]
budget_row_to_item = {it["legacy"]["budget_row"]: it["item_id"] for it in capex if it["legacy"].get("budget_row")}
cf_row_to_item = {}
for it in capex:
    for r in it["legacy"].get("cf_rows", []):
        cf_row_to_item.setdefault(r, it["item_id"])

# ---------------------------------------------------------------- значения
def value_target(sheet, row, col):
    c = gl(col)
    T = {
        "ТЭПы": {
            5: ("TIME.MILESTONES", "remove", "срок проекта вычисляется из вех"),
            13: ("LAND.CADASTRAL_VALUE", "keep" if c == "C" else "remove", "" if c == "C" else "дубль"),
            15: ("LAND.AREA", "keep" if c == "C" else "remove", "" if c == "C" else "дубль"),
            17: ("GEN.PHASES_COUNT", "fix", "очереди должны считаться раздельно"),
            19: ("TEP.GFA_ABOVE", "keep", ""),
            21: ("TEP.NONRES_GFA", "keep", ""),
            22: ("TEP.APT_AREA", "keep", ""),
            23: ("TEP.COMM_AREA", "remove", "коэффициент 0,8 без назначения"),
            27: ("TEP.PARKING_COUNT_OVERRIDE", "fix", "ниже норматива 2118-ПП"),
            28: ("TEP.APT_MIX", "remove", "дубль количества квартир"),
            29: ("TEP.LANDSCAPE_SHARE", "keep", "нужна ссылка на ППТ/НГП"),
            30: ("TEP.ROAD_SHARE", "keep", "нужна ссылка на ППТ"),
            31: ("TEP.GREEN_SHARE", "keep", "нужна ссылка на ППТ/НГП"),
            35: ("CAPEX.ITEMS", "replace", "база только для режима совместимости: перевод вбитых сумм бюджета в ставки на м²; продаваемая площадь — F.TEP.SALEABLE_AREA"),
            44: ("TEP.APT_MIX" if c == "F" else "TEP.PARKING_COUNT_OVERRIDE", "fix", "итог вбит числом"),
        },
    }
    if sheet == "ТЭПы":
        if row in T["ТЭПы"]:
            return T["ТЭПы"][row]
        if row in (41, 42, 43):
            return {
                "C": ("TEP.APT_MIX", "remove", "доля вводилась вручную — теперь F.TEP.APT_SHARE"),
                "D": ("TEP.APT_MIX", "keep", "avg_area"),
                "F": ("TEP.APT_MIX", "keep", "count"),
                "G": ("SALES.PRODUCTS", "keep", "start_price, источник не указан"),
                "I": ("TEP.PARKING_NORM", "keep", "совпадает с ПП Москвы 2118-ПП"),
            }.get(c, ("TEP.APT_MIX", "clarify", ""))
        if row == 48:
            return {"C": ("TEP.COMM_AREA", "remove", "10 332 vs 10 322 в C23 — опечатка"),
                    "D": ("SALES.PRODUCTS", "keep", "средняя площадь лота ПСН"),
                    "G": ("SALES.PRODUCTS", "keep", "цена ПСН руб/м²")}.get(c)
        if row == 49:
            return {"C": ("TEP.PARKING_AREA_PER_SPACE", "remove", "5 200 = 127 м/м × 40,945 м² — остаток прежней версии, в расчётах не участвует"),
                    "D": ("TEP.PARKING_AREA_PER_SPACE", "keep", ""),
                    "G": ("SALES.PRODUCTS", "fix", "цена м/м в руб/м² × 40,945")}.get(c)
        if row == 50:
            return ("TEP.STORAGE_COUNT", "keep", "прочие помещения = 0")
    if sheet == "Бюджет":
        special = {
            (19, "D"): ("LAND.AGENT_FEE_RATE", "fix", "ставка не участвует в сумме"),
            (22, "D"): ("LAND.VRI_FEE", "replace", "5 000 руб/м² без источника"),
            (48, "D"): ("OPEX.DEV_FEE_RATE", "fix", "сумма F48 вбита отдельно"),
            (51, "D"): ("OPEX.MARKETING_RATE", "keep", ""),
            (52, "D"): ("OPEX.BROKERAGE_RATE", "keep", ""),
            (54, "D"): ("TAX.VAT_RATE", "keep", ""),
            (55, "D"): ("TAX.PROFIT_RATE", "keep", ""),
            (65, "D"): ("FIN.EQUITY_SHARE", "remove", "дубль CF1!D132"),
            (68, "D"): ("F.FIN.FEES", "remove", "«банковские расходы 15%» дублируют проценты"),
            (21, "F"): ("LAND.RENT_ANNUAL", "replace", "итог без расчёта"),
            (24, "F"): ("LAND.CITY_OBLIGATIONS", "fix", "перенести в правообладание"),
        }
        if (row, c) in special:
            return special[(row, c)]
        if row in budget_row_to_item:
            return (f"CAPEX.ITEMS[{budget_row_to_item[row]}]", "replace", "сумма/ставка вбита числом → ставка × база")
    if sheet == "CF1":
        if row == 2: return ("GEN.MODEL_START_DATE", "keep", "")
        if row in (6, 7, 8): return ("TIME.MILESTONES", "remove", "флаги/номера проставлены вручную")
        if row in cf_row_to_item or (row - 1) in cf_row_to_item:
            item = cf_row_to_item.get(row) or cf_row_to_item.get(row - 1)
            return (f"CAPEX.ITEMS[{item}].schedule", "replace", "ручной % распределения → правило графика; ряд сохранён в кейсе как manual")
        special = {
            84: ("TAX.LAND_RATE", "replace", "0,2% без источника"),
            86: ("TAX.PROFIT_RATE", "remove", "дубль"),
            92: ("FIN.ESCROW_RESERVE_RATE", "keep", ""),
            102: ("FIN.COLLATERAL_DISCOUNT", "keep", ""),
            111: ("FIN.FEE_ARRANGEMENT", "keep", ""),
            115: ("FIN.KEY_RATE_PATH", "replace", "14,25% устарело: 14,00% на 11.09.2026"),
            116: ("FIN.RATE_PREFERENTIAL", "keep", ""),
            117: ("FIN.RATE_BASE_SPREAD", "fix", "база 20% → ключевая + спред"),
            118: ("FIN.RATE_DISCOUNT_COEF", "keep", ""),
            119: ("FIN.RATE_MIN", "keep", ""),
            132: ("FIN.EQUITY_SHARE", "fix", "взнос вперёд, а не % от расходов периода"),
            139: ("VAL.EQUITY_PREMIUM", "fix", "разложить на безрисковую + премию"),
        }
        if row == 85:
            return ("TAX.VAT_RATE", "remove", "дубль") if c == "D" else ("TAX.VAT_REGIME", "remove", "НДС 5 млн/кв вбит вручную")
        if row in special: return special[row]
    if sheet == "План продаж":
        if row in (1, 2, 3): return ("SALES.PACE", "remove", "дубль рядов 28/33/38")
        if row in (21, 22): return ("SALES.PACE", "remove", "итоги вбиты числами")
        if row in (28, 33, 38): return ("SALES.PACE", "keep", "шт/квартал")
        if row == 43: return ("SALES.PACE", "fix", "ПСН: дробные лоты, продано 10 888 м² > запас 10 322 м²")
        if row == 48: return ("SALES.PACE", "keep", "машино-места, шт/квартал")
        if row in (30, 35, 40, 45, 50, 55):
            if row == 55 and c != "E":
                return ("SALES.PRICE_STAGE_UPLIFT", "remove", "ошибка ввода 2.02…20.02")
            return ("SALES.PRICE_STAGE_UPLIFT", "replace", "2%/кв → рынок + стадия")
        if row == 62: return ("TAX.VAT_RATE", "remove", "дубль")
        if row == 70: return ("—", "remove", "пустой ряд без подписи")
    if sheet == "Эскроу":
        return ("SALES.PAYMENT_MIX", "keep" if row in (3, 5, 6) else "fix", "")
    return ("UNMAPPED", "clarify", "")

rows_out = []
text_out = []
for ws in wb.worksheets:
    sname = ws.title.strip()
    wsv = wv[ws.title]
    for row in ws.iter_rows(max_col=80):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            label = ""
            for col in (1, 2, 3):
                lv = ws.cell(cell.row, col).value
                if isinstance(lv, str) and not lv.startswith("="):
                    label = lv.strip()
            if isinstance(v, str) and v.startswith("="):
                continue
            if isinstance(v, (int, float)) or hasattr(v, "year"):
                tgt = value_target(sname, cell.row, cell.column) or ("UNMAPPED", "clarify", "")
                rows_out.append([sname, cell.coordinate, label, v, *tgt])
            else:
                text_out.append([sname, cell.coordinate, v])

with open(ROOT / "legacy" / "legacy_values_map.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["sheet", "cell", "row_label", "value", "target_id", "verdict", "note"])
    w.writerows(rows_out)
with open(ROOT / "legacy" / "legacy_text_cells.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["sheet", "cell", "text"])
    w.writerows(text_out)

# ---------------------------------------------------------------- формулы
ref = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\(\w])")
def rel(f, r, c):
    def sub(m):
        a, col, b, row = m.groups()
        cc, rr = ci(col), int(row)
        return (f"R{rr}" if b else f"R[{rr-r}]") + (f"C{cc}" if a else f"C[{cc-c}]")
    return ref.sub(sub, f)

def formula_target(sheet, row, col):
    c = gl(col)
    if sheet == "ТЭПы":
        m = {20: "TEP.RES_GFA", 21: "TEP.NONRES_GFA", 22: "TEP.APT_EFFICIENCY", 23: "TEP.COMM_AREA", 26: "TEP.MOP_AREA",
             27: "F.TEP.PARKING_COUNT", 28: "F.TEP.PARKING_COUNT", 29: "F.TEP.LANDSCAPE_AREA", 30: "F.TEP.LANDSCAPE_AREA",
             31: "F.TEP.LANDSCAPE_AREA", 32: "F.TEP.LANDSCAPE_AREA", 33: "F.TEP.LANDSCAPE_AREA", 34: "F.TEP.GFA_BELOW_EST",
             38: "F.TEP.SALEABLE_AREA", 44: "F.TEP.APT_TYPE_AREA", 51: "F.TEP.SALEABLE_AREA"}
        if row in (41, 42, 43):
            return {"E": ("F.TEP.APT_TYPE_AREA", "keep"), "H": ("F.SALES.WAVG_PRICE", "keep"), "J": ("F.TEP.PARKING_REQUIRED", "keep")}.get(c, ("F.TEP.APT_TYPE_AREA", "keep"))
        if row == 48: return ("F.TEP.SALEABLE_AREA", "keep")
        if row == 49:
            return {"E": ("F.TEP.GFA_BELOW_EST", "keep"), "F": ("F.TEP.PARKING_COUNT", "fix"), "H": ("SALES.PRODUCTS", "fix")}.get(c, ("F.TEP.GFA_BELOW_EST", "keep"))
        if row == 50: return ("F.TEP.SALEABLE_AREA", "keep")
        if row == 44 and c == "J": return ("F.TEP.PARKING_REQUIRED", "fix")
        if row in m: return (m[row], "keep")
    if sheet == "Бюджет":
        if 3 <= row <= 12: return ("F.SALES.REVENUE_TOTAL", "fix")
        if 16 <= row <= 52:
            if c == "G": return ("F.KPI.MARGIN", "keep")
            if row in (42,): return ("F.CAPEX.ITEM_TOTAL", "fix")
            return ("F.CAPEX.ITEM_TOTAL", "replace")
        m = {53: ("F.TAX.PAYMENTS", "replace"), 54: ("F.TAX.VAT_PAYABLE", "replace"), 55: ("F.TAX.PROFIT_TAX", "replace"),
             56: ("F.FIN.INTEREST", "keep"), 57: ("F.FIN.INTEREST", "keep"), 59: ("F.CAPEX.TOTAL", "fix"),
             61: ("F.KPI.MARGIN", "fix"), 63: ("F.KPI.NPV", "replace"), 64: ("F.KPI.IRR", "replace"),
             65: ("F.FIN.EQUITY_REQUIRED", "fix"), 66: ("F.KPI.IRR", "replace"), 67: ("F.FIN.LIMIT", "fix"),
             68: ("F.FIN.FEES", "remove"), 69: ("F.KPI.COST_PER_M2", "fix"), 70: ("F.SALES.WAVG_PRICE", "fix"),
             71: ("F.KPI.COST_PER_M2", "fix"), 73: ("F.KPI.UNFORECASTED_REVENUE", "replace"),
             74: ("F.KPI.UNFORECASTED_REVENUE", "replace"), 75: ("F.KPI.UNFORECASTED_REVENUE", "replace")}
        if row in m: return m[row]
    if sheet == "CF1":
        if row in (2, 3): return ("F.TIME.DATE", "keep")
        if row in (4, 5): return ("F.TIME.FLAG_CONSTRUCTION", "fix")
        if row == 6: return ("F.TIME.FLAG_ESCROW_RELEASE", "fix")
        if row == 7: return ("F.TIME.FLAG_CONSTRUCTION", "remove")
        if row in (8, 9): return ("F.TIME.FLAG_PRESALE", "fix")
        if row in (10, 11): return ("F.TEP.SALEABLE_AREA", "fix")
        if row == 15: return ("F.SALES.REVENUE_TOTAL", "fix")
        if row == 78: return ("F.CAPEX.ITEM_CASH", "fix")
        if row == 79: return ("F.CAPEX.ITEM_CASH", "fix")
        if row in (18, 19, 20, 23, 27, 40, 77): return ("F.CAPEX.TOTAL", "fix")
        if 21 <= row <= 76: return ("F.CAPEX.ITEM_CASH", "replace")
        if row in (81, 84, 85, 86): return ("F.TAX.PAYMENTS", "replace")
        if 90 <= row <= 95: return ("F.ESC.BALANCE", "fix")
        if row in (97, 98): return ("F.FIN.DRAW", "fix")
        if 99 <= row <= 106: return ("F.FIN.REPAYMENT", "fix")
        if 107 <= row <= 110: return ("F.FIN.INTEREST", "keep")
        if row in (111, 112): return ("F.FIN.FEES", "fix")
        if row == 113: return ("F.FIN.REPAYMENT", "keep")
        if 115 <= row <= 125: return ("F.FIN.RATE", "fix")
        if row in (127, 128): return ("F.FIN.EFFECTIVE_RATE", "fix")
        if row == 132: return ("F.FIN.EQUITY_IN", "replace")
        if row == 133: return ("F.CF.CFADS", "fix")
        if 134 <= row <= 137: return ("F.CF.CASH_BALANCE", "fix")
        if row == 138: return ("F.CF.FCFE", "replace")
        if 139 <= row <= 142: return ("F.KPI.NPV", "replace")
        if row == 143: return ("F.KPI.IRR", "replace")
        if row in (144, 145): return ("F.KPI.LTC_LTV", "replace")
        if 148 <= row <= 155: return ("F.KPI.LLCR", "replace")
    if sheet == "План продаж":
        if row in (1, 2, 3, 4, 21, 22): return ("SALES.PACE", "remove")
        if 5 <= row <= 18: return ("SALES.PRODUCTS", "keep")
        if row in (25, 26): return ("F.SALES.CONTRACT_VALUE", "keep")
        if 27 <= row <= 56:
            off = (row - 27) % 5
            return [("F.SALES.CONTRACT_VALUE", "keep"), ("SALES.PACE", "keep"), ("F.SALES.SOLD_AREA", "fix"),
                    ("SALES.PRICE_STAGE_UPLIFT", "replace"), ("F.SALES.PRICE", "replace")][off]
        if row == 58: return ("F.SALES.WAVG_PRICE", "replace")
        if row == 59: return ("F.SALES.SOLD_AREA", "keep")
        if row in (60, 61): return ("F.TAX.OUTPUT_VAT", "fix")
        if row == 65: return ("F.SALES.END_PRICE", "replace")
        if row == 66: return ("F.SALES.WAVG_PRICE", "replace")
        if row == 67: return ("F.SALES.SOLD_AREA", "keep")
        if row in (70, 71): return ("—", "remove")
    if sheet == "Эскроу":
        if row in (1, 2): return ("—", "remove")
        return ("F.SALES.CASH_IN", "replace")
    if sheet == "Dashboard":
        m = {**{r: ("F.TEP.SALEABLE_AREA", "fix") for r in range(2, 18)},
             19: ("F.SALES.REVENUE_TOTAL", "fix"), 21: ("F.SALES.PRICE", "keep"), 23: ("F.SALES.END_PRICE", "replace"),
             24: ("F.SALES.WAVG_PRICE", "replace"), 25: ("F.SALES.SOLD_AREA", "keep"), 26: ("F.SALES.SOLD_AREA", "keep"),
             **{r: ("F.CAPEX.TOTAL", "keep") for r in range(27, 33)},
             33: ("F.KPI.MARGIN", "keep"), 34: ("F.KPI.MARGIN", "keep"), 35: ("F.TAX.PAYMENTS", "fix"),
             36: ("F.FIN.LIMIT", "fix"), 37: ("F.FIN.INTEREST", "keep"), 38: ("F.FIN.DRAW", "fix"),
             39: ("F.FIN.EQUITY_IN", "fix"), 40: ("F.KPI.MARGIN", "fix"), 41: ("F.CF.CFADS", "fix"), 42: ("F.KPI.MARGIN", "fix"),
             43: ("F.KPI.MARGIN", "replace"), 44: ("F.KPI.IRR", "replace"), 45: ("F.KPI.IRR", "replace"),
             46: ("F.KPI.LTC_LTV", "replace"), 50: ("F.KPI.COST_PER_M2", "fix"), 51: ("F.KPI.LLCR", "replace"),
             52: ("F.KPI.LLCR", "replace"), 57: ("F.KPI.NPV", "replace"), 58: ("F.KPI.LLCR", "replace"),
             **{r: ("F.KPI.LLCR", "replace") for r in range(61, 66)}}
        if row in m: return m[row]
    return ("UNMAPPED", "clarify")

pat = {}
for ws in wb.worksheets:
    sname = ws.title.strip()
    for row in ws.iter_rows(max_col=80):
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            key = (sname, cell.row, rel(v, cell.row, cell.column))
            if key not in pat:
                label = ""
                for col in (1, 2, 3):
                    lv = ws.cell(cell.row, col).value
                    if isinstance(lv, str) and not lv.startswith("="):
                        label = lv.strip()
                pat[key] = {"first": cell.coordinate, "last": cell.coordinate, "n": 0, "sample": v, "label": label, "col": cell.column}
            pat[key]["last"] = cell.coordinate
            pat[key]["n"] += 1

with open(ROOT / "legacy" / "legacy_formulas_map.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["sheet", "row", "row_label", "range", "cells", "sample_formula", "r1c1_pattern", "target_id", "verdict"])
    for (s, r, p), d in pat.items():
        tgt = formula_target(s, r, d["col"])
        w.writerow([s, r, d["label"], f'{d["first"]}:{d["last"]}', d["n"], d["sample"], p, *tgt])

n_vals = len(rows_out); n_unm = sum(1 for r in rows_out if r[4] == "UNMAPPED")
n_f = len(pat)
n_fu = sum(1 for (s, r, p), d in pat.items() if formula_target(s, r, d["col"])[0] == "UNMAPPED")
n_fcells = sum(d["n"] for d in pat.values())
print(f"values: {n_vals} (unmapped {n_unm}); text cells: {len(text_out)}; formula patterns: {n_f} covering {n_fcells} cells (unmapped {n_fu})")
if n_unm or n_fu:
    sys.exit(1)
